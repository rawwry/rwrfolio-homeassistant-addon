import os
import io
import csv
import datetime
from fastapi import FastAPI, Request, Form, UploadFile, File, Depends, HTTPException, status
from fastapi.responses import HTMLResponse, RedirectResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
import openpyxl

from database import get_connection, init_db
from calculations import get_portfolio_summary, get_coin_display_name
from importer import parse_uploaded_file, clean_float
from price_service import update_live_prices

# Initialize FastAPI App
app = FastAPI(title="Crypto Tracker - Synology NAS & Home Assistant Edition", version="1.0.0")

# Setup Paths & Jinja2 Templates
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
STATIC_DIR = os.path.join(BASE_DIR, "static")

os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
templates = Jinja2Templates(directory=TEMPLATES_DIR)

# Custom Jinja2 Filters for Euro Formatting
def format_currency_eur(value: float, decimals: int = 2) -> str:
    try:
        val = float(value or 0.0)
        formatted = f"{val:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{formatted} €"
    except Exception:
        return f"{value} €"

def format_crypto_amount(value: float) -> str:
    try:
        val = float(value or 0.0)
        if abs(val) >= 1000:
            return f"{val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        elif abs(val) >= 1:
            return f"{val:,.4f}".replace(",", "X").replace(".", ",").replace("X", ".")
        else:
            return f"{val:.8f}".rstrip('0').rstrip('.')
    except Exception:
        return str(value)

def format_percentage(value: float) -> str:
    try:
        val = float(value or 0.0)
        sign = "+" if val > 0 else ""
        return f"{sign}{val:.2f}%"
    except Exception:
        return f"{value}%"

templates.env.filters["eur"] = format_currency_eur
templates.env.filters["crypto"] = format_crypto_amount
templates.env.filters["pct"] = format_percentage

@app.on_event("startup")
def on_startup():
    init_db()

# Dependency for DB connection per request
def get_db():
    conn = get_connection()
    try:
        yield conn
    finally:
        conn.close()

# ----------------- ROUTES ----------------- #

@app.get("/", response_class=HTMLResponse)
async def route_dashboard(request: Request, db=Depends(get_db)):
    summary = get_portfolio_summary(db)
    
    # Get last 8 recent transactions
    cursor = db.cursor()
    cursor.execute("""
        SELECT * FROM transactions 
        ORDER BY timestamp DESC, id DESC 
        LIMIT 8;
    """)
    recent_txs = [dict(row) for row in cursor.fetchall()]
    
    return templates.TemplateResponse("dashboard.html", {
        "request": request,
        "active_page": "dashboard",
        "totals": summary["totals"],
        "assets": summary["assets"],
        "recent_transactions": recent_txs,
        "now": datetime.datetime.utcnow().strftime("%d.%m.%Y %H:%M UTC")
    })

@app.get("/transactions", response_class=HTMLResponse)
async def route_transactions(
    request: Request,
    asset: str = "ALL",
    type_filter: str = "ALL",
    search: str = "",
    page: int = 1,
    db=Depends(get_db)
):
    cursor = db.cursor()
    
    # Query all available coins for the filter dropdown
    cursor.execute("""
        SELECT DISTINCT received_currency FROM transactions WHERE received_currency IS NOT NULL AND received_currency != ''
        UNION
        SELECT DISTINCT spent_currency FROM transactions WHERE spent_currency IS NOT NULL AND spent_currency != '' AND spent_currency != 'EUR';
    """)
    available_assets = sorted([r[0] for r in cursor.fetchall() if r[0]])

    # Build filtered query
    query = "SELECT * FROM transactions WHERE 1=1"
    params = []

    if asset != "ALL" and asset.strip():
        query += " AND (received_currency = ? OR spent_currency = ?)"
        params.extend([asset.strip().upper(), asset.strip().upper()])

    if type_filter != "ALL" and type_filter.strip():
        query += " AND type = ?"
        params.append(type_filter.strip().upper())

    if search.strip():
        search_pattern = f"%{search.strip()}%"
        query += " AND (description LIKE ? OR id LIKE ? OR notes LIKE ? OR transaction_kind LIKE ?)"
        params.extend([search_pattern, search_pattern, search_pattern, search_pattern])

    # Count total matching rows
    count_query = query.replace("SELECT *", "SELECT COUNT(*) as count")
    cursor.execute(count_query, params)
    total_count = cursor.fetchone()["count"]

    # Pagination: 25 items per page
    limit = 25
    offset = (page - 1) * limit
    query += " ORDER BY timestamp DESC, id DESC LIMIT ? OFFSET ?"
    params.extend([limit, offset])

    cursor.execute(query, params)
    transactions = [dict(row) for row in cursor.fetchall()]
    total_pages = max(1, (total_count + limit - 1) // limit)

    summary = get_portfolio_summary(db)

    return templates.TemplateResponse("transactions.html", {
        "request": request,
        "active_page": "transactions",
        "transactions": transactions,
        "total_count": total_count,
        "page": page,
        "total_pages": total_pages,
        "current_asset": asset,
        "current_type": type_filter,
        "current_search": search,
        "available_assets": available_assets,
        "totals": summary["totals"]
    })

@app.post("/transactions/add")
async def route_add_transaction(
    timestamp: str = Form(...),
    type: str = Form(...),
    source: str = Form("manual"),
    description: str = Form("Manueller Eintrag"),
    spent_currency: str = Form("EUR"),
    spent_amount: float = Form(0.0),
    received_currency: str = Form(""),
    received_amount: float = Form(0.0),
    price_per_unit_eur: float = Form(0.0),
    notes: str = Form(""),
    db=Depends(get_db)
):
    if price_per_unit_eur == 0.0 and received_amount > 0 and spent_amount > 0:
        price_per_unit_eur = spent_amount / received_amount

    tx_id = f"man_{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{received_currency}"
    cursor = db.cursor()
    cursor.execute("""
        INSERT INTO transactions (
            id, timestamp, source, type, description, spent_currency, spent_amount,
            received_currency, received_amount, price_per_unit_eur, native_currency,
            native_amount, transaction_kind, fee, fee_currency, notes, imported_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        tx_id, timestamp, source, type.upper(), description,
        spent_currency.upper(), spent_amount, received_currency.upper(), received_amount,
        price_per_unit_eur, "EUR", spent_amount, "manual_entry", 0.0, "", notes,
        datetime.datetime.utcnow().isoformat()
    ))
    db.commit()

    return RedirectResponse(url="/transactions?msg=created", status_code=status.HTTP_303_SEE_OTHER)

@app.post("/transactions/{tx_id}/delete")
async def route_delete_transaction(tx_id: str, db=Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("DELETE FROM transactions WHERE id = ?;", (tx_id,))
    db.commit()
    return RedirectResponse(url="/transactions?msg=deleted", status_code=status.HTTP_303_SEE_OTHER)

@app.post("/transactions/bulk-delete")
async def route_bulk_delete(ids: list[str] = Form(...), db=Depends(get_db)):
    if ids:
        cursor = db.cursor()
        placeholders = ",".join(["?"] * len(ids))
        cursor.execute(f"DELETE FROM transactions WHERE id IN ({placeholders});", ids)
        db.commit()
    return RedirectResponse(url="/transactions?msg=bulk_deleted", status_code=status.HTTP_303_SEE_OTHER)

@app.get("/import", response_class=HTMLResponse)
async def route_import_view(request: Request, db=Depends(get_db)):
    summary = get_portfolio_summary(db)
    return templates.TemplateResponse("import.html", {
        "request": request,
        "active_page": "import",
        "totals": summary["totals"]
    })

@app.post("/import/upload", response_class=HTMLResponse)
async def route_import_upload(
    request: Request,
    file: UploadFile = File(...),
    db=Depends(get_db)
):
    content = await file.read()
    transactions, detected_source, errors = parse_uploaded_file(content, file.filename or "export.csv")

    imported_count = 0
    skipped_duplicates = 0

    if transactions:
        cursor = db.cursor()
        for tx in transactions:
            # Check for existing duplicate by id or identical timestamp + coin + amount
            cursor.execute("""
                SELECT id FROM transactions 
                WHERE id = ? OR (timestamp = ? AND received_currency = ? AND received_amount = ?);
            """, (tx["id"], tx["timestamp"], tx["received_currency"], tx["received_amount"]))
            
            if cursor.fetchone():
                skipped_duplicates += 1
            else:
                cursor.execute("""
                    INSERT INTO transactions (
                        id, timestamp, source, type, description, spent_currency, spent_amount,
                        received_currency, received_amount, price_per_unit_eur, native_currency,
                        native_amount, transaction_kind, fee, fee_currency, notes, imported_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    tx["id"], tx["timestamp"], tx["source"], tx["type"], tx["description"],
                    tx["spent_currency"], tx["spent_amount"], tx["received_currency"], tx["received_amount"],
                    tx["price_per_unit_eur"], tx["native_currency"], tx["native_amount"], tx["transaction_kind"],
                    tx["fee"], tx["fee_currency"], tx["notes"], tx["imported_at"]
                ))
                imported_count += 1
        db.commit()

    summary = get_portfolio_summary(db)

    return templates.TemplateResponse("import.html", {
        "request": request,
        "active_page": "import",
        "totals": summary["totals"],
        "success_msg": f"{imported_count} Transaktionen erfolgreich importiert ({skipped_duplicates} Duplikate übersprungen).",
        "errors": errors,
        "detected_source": detected_source,
        "imported_count": imported_count,
        "skipped_count": skipped_duplicates
    })

@app.post("/prices/update")
async def route_update_prices(
    symbol: str = Form(None),
    price_eur: float = Form(None),
    action: str = Form("refresh_all"),
    db=Depends(get_db)
):
    if action == "refresh_all":
        summary = get_portfolio_summary(db)
        symbols = [a["symbol"] for a in summary["assets"]]
        update_live_prices(db, symbols)
    elif symbol and price_eur is not None and price_eur > 0:
        cursor = db.cursor()
        now_str = datetime.datetime.utcnow().isoformat()
        cursor.execute("""
            INSERT INTO custom_prices (symbol, price_eur, source, updated_at)
            VALUES (?, ?, 'manual', ?)
            ON CONFLICT(symbol) DO UPDATE SET
                price_eur = excluded.price_eur,
                source = 'manual',
                updated_at = excluded.updated_at;
        """, (symbol.upper(), price_eur, now_str))
        db.commit()

    return RedirectResponse(url="/?msg=prices_updated", status_code=status.HTTP_303_SEE_OTHER)

@app.get("/export/csv")
async def route_export_csv(db=Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("SELECT * FROM transactions ORDER BY timestamp ASC;")
    rows = cursor.fetchall()

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    
    # Headers
    writer.writerow([
        "ID", "Timestamp (UTC)", "Source", "Type", "Description", 
        "Spent Currency", "Spent Amount", "Received Currency", "Received Amount", 
        "Price Per Unit (EUR)", "Native Currency", "Native Amount", "Transaction Kind", "Notes"
    ])

    for r in rows:
        writer.writerow([
            r["id"], r["timestamp"], r["source"], r["type"], r["description"],
            r["spent_currency"], r["spent_amount"], r["received_currency"], r["received_amount"],
            r["price_per_unit_eur"], r["native_currency"], r["native_amount"], r["transaction_kind"], r["notes"]
        ])

    csv_data = output.getvalue()
    filename = f"crypto_transactions_{datetime.datetime.utcnow().strftime('%Y%m%d')}.csv"
    
    return Response(
        content=csv_data,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/export/xlsx")
async def route_export_xlsx(db=Depends(get_db)):
    cursor = db.cursor()
    cursor.execute("SELECT * FROM transactions ORDER BY timestamp ASC;")
    rows = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Crypto Transactions"

    headers = [
        "ID", "Timestamp (UTC)", "Source", "Type", "Description", 
        "Spent Currency", "Spent Amount", "Received Currency", "Received Amount", 
        "Price Per Unit (EUR)", "Native Currency", "Native Amount", "Notes"
    ]
    ws.append(headers)

    for r in rows:
        ws.append([
            r["id"], r["timestamp"], r["source"], r["type"], r["description"],
            r["spent_currency"], r["spent_amount"], r["received_currency"], r["received_amount"],
            r["price_per_unit_eur"], r["native_currency"], r["native_amount"], r["notes"]
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"crypto_transactions_{datetime.datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/homeassistant", response_class=HTMLResponse)
async def route_homeassistant(request: Request, db=Depends(get_db)):
    summary = get_portfolio_summary(db)
    return templates.TemplateResponse("homeassistant.html", {
        "request": request,
        "active_page": "homeassistant",
        "totals": summary["totals"],
        "assets": summary["assets"],
        "host_header": request.headers.get("host", "synology-ip:8000")
    })

# Home Assistant REST Sensor API Endpoint
@app.get("/api/homeassistant/sensors")
async def route_ha_sensors(db=Depends(get_db)):
    """
    Returns structured JSON formatted specifically for Home Assistant REST integration.
    """
    summary = get_portfolio_summary(db)
    totals = summary["totals"]
    assets = summary["assets"]

    asset_dict = {}
    for a in assets:
        asset_dict[a["symbol"]] = {
            "balance": a["current_balance"],
            "value_eur": round(a["current_value_eur"], 2),
            "invested_eur": round(a["total_invested_eur"], 2),
            "avg_buy_price_eur": round(a["average_buy_price"], 4),
            "current_price_eur": round(a["current_price"], 4),
            "pnl_eur": round(a["pnl_eur"], 2),
            "pnl_percentage": round(a["pnl_percentage"], 2),
            "allocation_percentage": round(a["allocation_percentage"], 2)
        }

    return JSONResponse({
        "status": "online",
        "timestamp": datetime.datetime.utcnow().isoformat(),
        "portfolio": {
            "total_value_eur": round(totals["current_value_eur"], 2),
            "total_invested_eur": round(totals["total_invested_eur"], 2),
            "total_pnl_eur": round(totals["total_pnl_eur"], 2),
            "total_pnl_percentage": round(totals["total_pnl_percentage"], 2),
            "asset_count": totals["asset_count"],
            "total_transactions": totals["total_tx_count"],
            "top_asset": totals["top_asset"],
            "top_asset_share": round(totals["top_asset_share"], 2)
        },
        "assets": asset_dict
    })
